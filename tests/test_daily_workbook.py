"""
### FILE: tests/test_daily_workbook.py
Edge-case tests for feeder_forecast/parser.py (spec section 12), built as
mutated copies of the real sample workbook
(Claude outputs/SHIRORO_05102026.xlsx) wherever row/cell-level behaviour
is under test, and bare Control-sheet-only workbooks for the region/date
resolution matrix, which never touches Record and Accounting at all.
"""
import datetime as dt
import shutil

import openpyxl
import pytest
from sqlalchemy import text

from feeder_forecast.parser import parse_control_sheet, parse_workbook

TODAY = dt.date.today()


def _control_only_wb(region_cell=None, date_cell=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Control"
    ws["L3"] = region_cell
    ws["L4"] = date_cell
    return wb


# ---------------------------------------------------------------------
# Region/date resolution matrix (spec section 3a)
# ---------------------------------------------------------------------

def test_region_date_both_match_resolves():
    wb = _control_only_wb("SHIRORO", dt.datetime(2026, 10, 5))
    r = parse_control_sheet(wb, "SHIRORO_05102026.xlsx")
    assert r.status == "resolved"
    assert r.region == "Shiroro" and r.date == dt.date(2026, 10, 5)
    assert r.region_source == "sheet" and r.date_source == "sheet"


def test_region_mismatch_pauses_r05():
    wb = _control_only_wb("SHIRORO", dt.datetime(2026, 10, 5))
    r = parse_control_sheet(wb, "ABUJA_05102026.xlsx")
    assert r.status == "pause"
    assert r.pause_code == "R05"
    assert r.pause_on in ("region", "both")
    # the date side matched, so it should already be resolved through
    assert r.date == dt.date(2026, 10, 5)


def test_date_mismatch_pauses_r05():
    wb = _control_only_wb("SHIRORO", dt.datetime(2026, 10, 5))
    r = parse_control_sheet(wb, "SHIRORO_06102026.xlsx")
    assert r.status == "pause"
    assert r.pause_code == "R05"
    assert r.pause_on in ("date", "both")
    # the region side matched, so it should already be resolved through
    assert r.region == "Shiroro"


def test_blank_date_with_filename_date_pauses_r13():
    wb = _control_only_wb("SHIRORO", None)
    r = parse_control_sheet(wb, "SHIRORO_05102026.xlsx")
    assert r.status == "pause"
    assert r.pause_code == "R13"
    assert r.filename_date == dt.date(2026, 10, 5)
    # region resolved cleanly and is carried through
    assert r.region == "Shiroro" and r.region_source == "sheet"


def test_blank_region_with_filename_region_resolves_with_r12_warning():
    wb = _control_only_wb(None, dt.datetime(2026, 10, 5))
    r = parse_control_sheet(wb, "SHIRORO_05102026.xlsx")
    assert r.status == "resolved"
    assert r.region == "Shiroro" and r.region_source == "filename"
    assert any(w["rule_code"] == "R12" for w in r.warnings)


def test_future_date_rejects_r04():
    future = TODAY + dt.timedelta(days=5)
    wb = _control_only_wb("SHIRORO", dt.datetime(future.year, future.month, future.day))
    r = parse_control_sheet(wb, f"SHIRORO_{future:%d%m%Y}.xlsx")
    assert r.status == "reject"
    assert r.reject_code == "R04"


def test_blank_date_and_unreadable_filename_date_rejects_r04():
    wb = _control_only_wb("SHIRORO", None)
    r = parse_control_sheet(wb, "SHIRORO_daily.xlsx")
    assert r.status == "reject"
    assert r.reject_code == "R04"


def test_invalid_region_text_rejects_r10():
    wb = _control_only_wb("NOTAREALREGION", dt.datetime(2026, 10, 5))
    r = parse_control_sheet(wb, "SHIRORO_05102026.xlsx")
    assert r.status == "reject"
    assert r.reject_code == "R10"


def test_blank_region_and_unreadable_filename_region_rejects_r10():
    wb = _control_only_wb(None, dt.datetime(2026, 10, 5))
    r = parse_control_sheet(wb, "05102026.xlsx")
    assert r.status == "reject"
    assert r.reject_code == "R10"


# ---------------------------------------------------------------------
# Row/cell-level rules (spec section 3b) -- real sample file, mutated
# ---------------------------------------------------------------------

@pytest.fixture
def workbook_copy(tmp_path, sample_path):
    dest = tmp_path / sample_path.name
    shutil.copy(sample_path, dest)
    return dest


def test_clean_sample_parses_ok(workbook_copy, engine):
    r = parse_workbook(workbook_copy, "Shiroro", dt.date(2026, 10, 5), engine)
    assert r.status == "ok"
    assert r.feeders_read == 59
    assert r.rows_saved == 59
    assert r.rows_skipped == 0


def test_layout_changed_rejects_r06(workbook_copy, engine):
    wb = openpyxl.load_workbook(workbook_copy)
    ws = wb["Record and Accounting"]
    ws.cell(2, 13).value = "WRONG"  # hour-1 label corrupted
    wb.save(workbook_copy)
    r = parse_workbook(workbook_copy, "Shiroro", dt.date(2026, 10, 5), engine)
    assert r.status == "reject"
    assert r.reject_code == "R06"


def test_text_in_forecast_cell_warns_r30(workbook_copy, engine):
    wb = openpyxl.load_workbook(workbook_copy)
    ws = wb["Record and Accounting"]
    ws.cell(4, 13).value = "OFF"  # row 4's hour-1 forecast cell
    wb.save(workbook_copy)
    r = parse_workbook(workbook_copy, "Shiroro", dt.date(2026, 10, 5), engine)
    assert r.status == "ok"
    assert r.rows_saved == 59  # row still saved
    r30 = [i for i in r.issues if i["rule_code"] == "R30" and i["cell"] == "M4"]
    assert len(r30) == 1
    hour1 = next(f for f in r.forecast_rows if f["feeder_id"] == r.daily_rows[0]["feeder_id"] and f["source_label"] == 1)
    assert hour1["forecast_mw"] is None


def test_duplicate_feeder_row_skipped_r21(workbook_copy, engine):
    wb = openpyxl.load_workbook(workbook_copy)
    ws = wb["Record and Accounting"]
    # make row 5 an identity-duplicate of row 4 (station + feeder name)
    ws.cell(5, 4).value = ws.cell(4, 4).value
    ws.cell(5, 8).value = ws.cell(4, 8).value
    wb.save(workbook_copy)
    r = parse_workbook(workbook_copy, "Shiroro", dt.date(2026, 10, 5), engine)
    assert r.status == "ok"
    assert r.rows_saved == 58  # one less than the clean 59
    assert r.rows_skipped == 1
    r21 = [i for i in r.issues if i["rule_code"] == "R21"]
    assert len(r21) == 1


def test_negative_forecast_warns_r31(workbook_copy, engine):
    wb = openpyxl.load_workbook(workbook_copy)
    ws = wb["Record and Accounting"]
    ws.cell(4, 13).value = -5
    wb.save(workbook_copy)
    r = parse_workbook(workbook_copy, "Shiroro", dt.date(2026, 10, 5), engine)
    assert r.status == "ok"
    r31 = [i for i in r.issues if i["rule_code"] == "R31"]
    assert len(r31) == 1


def test_decreasing_reading_warns_r32(workbook_copy, engine):
    wb = openpyxl.load_workbook(workbook_copy)
    ws = wb["Record and Accounting"]
    opening = ws.cell(4, 12).value
    ws.cell(4, 14).value = opening - 10  # hour-1 reading below opening
    wb.save(workbook_copy)
    r = parse_workbook(workbook_copy, "Shiroro", dt.date(2026, 10, 5), engine)
    assert r.status == "ok"
    r32 = [i for i in r.issues if i["rule_code"] == "R32"]
    assert len(r32) == 1


# ---------------------------------------------------------------------
# R14 / R35 -- needs previously-saved rows to compare against
# ---------------------------------------------------------------------

def test_r14_duplicate_file_pauses(workbook_copy, sample_path, engine):
    from feeder_forecast import store as ff_store

    base = parse_workbook(sample_path, "Shiroro", dt.date(2026, 10, 5), engine)
    upload_id = ff_store.save_upload(engine, base, sample_path.name, "pytest")
    try:
        again = parse_workbook(workbook_copy, "Shiroro", dt.date(2026, 10, 10), engine)
        assert again.status == "pause"
        assert again.pause_code == "R14"
        assert "identical" in again.pause_message
    finally:
        with engine.begin() as con:
            con.execute(text("DELETE FROM feeder_daily WHERE upload_id=:u"), dict(u=upload_id))
            con.execute(text("DELETE FROM feeder_hourly_forecast WHERE upload_id=:u"), dict(u=upload_id))
            con.execute(text("DELETE FROM feeder_meter_reading WHERE upload_id=:u"), dict(u=upload_id))
            con.execute(text("DELETE FROM daily_workbook_upload WHERE upload_id=:u"), dict(u=upload_id))
